"""Membuat Kalkulator-SUS-dan-Rekap-UAT.xlsx.

    python build_kalkulator.py <output.xlsx>
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
DATA = json.loads((Path(__file__).parent / "data_uat_sus.json").read_text())
N_RESP = 100
FIRST = 6                      # baris data pertama di sheet SUS (baris 5 = contoh)
LAST = FIRST + N_RESP - 1
GROUPS = ["Admin Mitra", "Calon Peserta"]

F = "Arial"
font = lambda **k: Font(name=F, size=k.pop("size", 10), **k)
INPUT = PatternFill("solid", fgColor="FFF59D")
HEAD = PatternFill("solid", fgColor="DCEFE4")
EXAMPLE = PatternFill("solid", fgColor="EEEEEE")
thin = Side(style="thin", color="8A949E")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)


def head(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.alignment, c.border = font(bold=True), HEAD, CENTER, BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


wb = Workbook()

# ------------------------------------------------------------ Petunjuk
ws = wb.active
ws.title = "Petunjuk"
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110
rows = [
    ("Kalkulator SUS & Rekap UAT", dict(bold=True, size=14)),
    ("Sistem Promosi Digital Terintegrasi Berbasis SEO dan Machine Learning — Politeknik Negeri Medan × Yayasan Sinergi Inspirasi Indonesia", dict(size=10, italic=True)),
    ("", {}),
    ("Cara pakai", dict(bold=True, size=11)),
    ("Sel berwarna KUNING adalah tempat mengisi data. Sel lain berisi rumus dan akan terhitung otomatis — jangan diubah.", {}),
    ("Sheet SUS: satu baris per responden. Pilih Kelompok (Admin Mitra / Calon Peserta) lalu isi jawaban butir 1–10 dengan angka 1–5 sesuai kuesioner.", {}),
    ("Baris 5 (abu-abu, kelompok \"Contoh\") hanya contoh pengisian dan tidak ikut dihitung.", {}),
    ("Skor SUS per responden muncul otomatis bila 10 butir terisi lengkap. Hasil per kelompok ada di sheet Ringkasan SUS.", {}),
    ("Sheet UAT: pilih Sesuai / Tidak Sesuai pada kolom Hasil untuk setiap skenario, dan tulis catatan bila perlu. Kesimpulan muncul otomatis di bagian atas.", {}),
    ("", {}),
    ("Rumus SUS (Brooke, 1996)", dict(bold=True, size=11)),
    ("Butir ganjil: jawaban − 1. Butir genap: 5 − jawaban. Jumlah kontribusi (0–40) × 2,5 = skor 0–100.", {}),
    ("Interpretasi penerimaan (Bangor dkk., 2008): ≥ 70 Acceptable, 50–69,9 Marginal, < 50 Not acceptable. Rata-rata umum SUS = 68.", {}),
    ("Kategori adjektif (disederhanakan dari Bangor dkk., 2009): > 85 Excellent–Best imaginable, 70–85 Good–Excellent, 50–69,9 OK–Good, < 50 Poor–Worst imaginable.", {}),
    ("Butir kuesioner: adaptasi Bahasa Indonesia mengacu pada " + DATA["sus_ref"] + ".", {}),
    ("", {}),
    ("Kriteria penerimaan UAT", dict(bold=True, size=11)),
    ("Diterima: ≥ 90% skenario Sesuai dan seluruh skenario kritis Sesuai. Diterima dengan perbaikan: ≥ 90% Sesuai, ada skenario non-kritis Tidak Sesuai. Belum diterima: < 90% Sesuai atau ada skenario kritis Tidak Sesuai.", {}),
]
for i, (text, style) in enumerate(rows, 1):
    c = ws.cell(row=i, column=2, value=text)
    c.font = font(**style)
    c.alignment = WRAP
ws.cell(row=5, column=2).fill = INPUT

# ------------------------------------------------------------ SUS
sus = wb.create_sheet("SUS")
sus["A1"] = "Data Jawaban Kuesioner SUS"
sus["A1"].font = font(bold=True, size=13)
sus["A2"] = "Isi sel kuning. Jawaban 1 = Sangat Tidak Setuju … 5 = Sangat Setuju."
sus["A2"].font = font(italic=True)
items = DATA["sus_items"]
cols = ["No", "Kelompok", "Nama / inisial"] + [f"B{i}" for i in range(1, 11)] + ["Skor SUS", "bantu"]
head(sus, 4, cols, [5, 16, 20] + [6] * 10 + [10, 6])
for i, text in enumerate(items):
    sus.cell(row=4, column=4 + i).comment = Comment(f"Butir {i + 1}: {text}", "Tim PTI")
sus.column_dimensions["O"].hidden = True   # kolom bantu untuk simpangan baku

def score_formula(r):
    q = lambda k: f"{chr(ord('D') + k - 1)}{r}"
    odd = ",".join(q(k) for k in (1, 3, 5, 7, 9))
    even = ",".join(q(k) for k in (2, 4, 6, 8, 10))
    return f'=IF(COUNT(D{r}:M{r})=10,(SUM({odd})-5+25-SUM({even}))*2.5,"")'

example = [4, 2, 5, 1, 4, 2, 4, 2, 5, 3]   # skor 80 (sama dengan contoh di kuesioner)
sus.cell(row=5, column=1, value="Contoh")
sus.cell(row=5, column=2, value="Contoh")
sus.cell(row=5, column=3, value="(contoh pengisian)")
for k, v in enumerate(example):
    sus.cell(row=5, column=4 + k, value=v)
sus.cell(row=5, column=14, value=score_formula(5))
for c in range(1, 16):
    cell = sus.cell(row=5, column=c)
    cell.fill, cell.font, cell.border, cell.alignment = EXAMPLE, font(italic=True, color="555555"), BOX, CENTER

for r in range(FIRST, LAST + 1):
    sus.cell(row=r, column=1, value=r - FIRST + 1)
    for c in range(2, 14):
        sus.cell(row=r, column=c).fill = INPUT
    sus.cell(row=r, column=14, value=score_formula(r)).number_format = "0.0"
    sus.cell(row=r, column=15, value=f'=IF(ISNUMBER(N{r}),N{r},0)')
    for c in range(1, 16):
        sus.cell(row=r, column=c).border = BOX
        sus.cell(row=r, column=c).alignment = CENTER
        if c != 1:
            sus.cell(row=r, column=c).font = font()
sus.freeze_panes = "D6"

dv_group = DataValidation(type="list", formula1='"' + ",".join(GROUPS) + '"', allow_blank=True,
                          error="Pilih Admin Mitra atau Calon Peserta", errorTitle="Kelompok")
dv_ans = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                        error="Isi angka 1 sampai 5", errorTitle="Jawaban SUS")
sus.add_data_validation(dv_group)
sus.add_data_validation(dv_ans)
dv_group.add(f"B{FIRST}:B{LAST}")
dv_ans.add(f"D{FIRST}:M{LAST}")

# ------------------------------------------------------------ Ringkasan SUS
rk = wb.create_sheet("Ringkasan SUS")
rk["A1"] = "Ringkasan Skor SUS per Kelompok"
rk["A1"].font = font(bold=True, size=13)
head(rk, 3, ["Ukuran"] + GROUPS, [34, 20, 20])
rng = lambda col: f"SUS!${col}${FIRST}:${col}${LAST}"
labels = ["Jumlah responden (lengkap)", "Rata-rata skor SUS", "Simpangan baku", "Skor terendah", "Skor tertinggi",
          "Penerimaan (Bangor dkk., 2008)", "Kategori adjektif", "Dibanding rata-rata umum (68)"]
for i, lab in enumerate(labels):
    r = 4 + i
    rk.cell(row=r, column=1, value=lab).font = font(bold=True)
    for j, g in enumerate(GROUPS):
        col = chr(ord("B") + j)
        n, mean = f"{col}$4", f"{col}$5"
        f = {
            0: f'=SUMPRODUCT(({rng("B")}="{g}")*ISNUMBER({rng("N")}))',
            1: f'=IF({n}=0,"-",AVERAGEIFS({rng("N")},{rng("B")},"{g}"))',
            2: f'=IF({n}<2,"-",SQRT(SUMPRODUCT(({rng("B")}="{g}")*ISNUMBER({rng("N")})*({rng("O")}-{mean})^2)/({n}-1)))',
            3: f'=IF({n}=0,"-",_xlfn.MINIFS({rng("N")},{rng("B")},"{g}"))',
            4: f'=IF({n}=0,"-",_xlfn.MAXIFS({rng("N")},{rng("B")},"{g}"))',
            5: f'=IF({n}=0,"-",IF({mean}>=70,"Acceptable",IF({mean}>=50,"Marginal","Not acceptable")))',
            6: f'=IF({n}=0,"-",IF({mean}>85,"Excellent – Best imaginable",IF({mean}>=70,"Good – Excellent",IF({mean}>=50,"OK – Good","Poor – Worst imaginable"))))',
            7: f'=IF({n}=0,"-",IF({mean}>=68,"Di atas rata-rata","Di bawah rata-rata"))',
        }[i]
        c = rk.cell(row=r, column=2 + j, value=f)
        c.alignment, c.border, c.font = CENTER, BOX, font()
        if i in (1, 2):
            c.number_format = "0.0"
    rk.cell(row=r, column=1).border = BOX

rk["A14"] = "Rata-rata jawaban per butir (1–5)"
rk["A14"].font = font(bold=True, size=11)
head(rk, 15, ["Butir"] + GROUPS)
for i, text in enumerate(items):
    r = 16 + i
    rk.cell(row=r, column=1, value=f"{i + 1}. {text}").alignment = WRAP
    rk.cell(row=r, column=1).border = BOX
    rk.cell(row=r, column=1).font = font(size=9)
    col_letter = chr(ord("D") + i)
    for j, g in enumerate(GROUPS):
        n = f"{chr(ord('B') + j)}$4"
        c = rk.cell(row=r, column=2 + j,
                    value=f'=IF({n}=0,"-",AVERAGEIFS({rng(col_letter)},{rng("B")},"{g}",{rng("N")},">=0"))')
        c.number_format, c.alignment, c.border, c.font = "0.00", CENTER, BOX, font()
rk.cell(row=27, column=1, value="Catatan: butir genap bernada negatif, sehingga nilai rendah pada butir genap justru baik.").font = font(italic=True, size=9)

# ------------------------------------------------------------ UAT
ut = wb.create_sheet("UAT")
ut["A1"] = "Rekap User Acceptance Test"
ut["A1"].font = font(bold=True, size=13)
scen = [(i, g) for g in DATA["uat"] for i in g["items"]]
S0 = 16                                    # baris skenario pertama (di bawah ringkasan)
S1 = S0 + len(scen) - 1
res, krit, grp = f"$E${S0}:$E${S1}", f"$D${S0}:$D${S1}", f"$B${S0}:$B${S1}"

head(ut, 3, ["Kelompok", "Jumlah", "Sesuai", "Tidak sesuai", "Belum diuji", "% Sesuai"],
     [28, 12, 12, 13, 12, 12])
ut.column_dimensions["C"].width = 46
ut.column_dimensions["F"].width = 34
for k, g in enumerate(DATA["uat"]):
    r = 4 + k
    name = f'{g["code"]}. {g["title"]}'
    vals = [name, f'=COUNTIF({grp},A{r})', f'=COUNTIFS({grp},A{r},{res},"Sesuai")',
            f'=COUNTIFS({grp},A{r},{res},"Tidak Sesuai")', f"=B{r}-C{r}-D{r}", f'=IF(B{r}=0,0,C{r}/B{r})']
    for c, v in enumerate(vals, 1):
        cell = ut.cell(row=r, column=c, value=v)
        cell.border, cell.font = BOX, font()
        cell.alignment = CENTER if c > 1 else WRAP
    ut.cell(row=r, column=6).number_format = "0.0%"
tr = 4 + len(DATA["uat"])
for c, v in enumerate(["Total", f"=SUM(B4:B{tr - 1})", f"=SUM(C4:C{tr - 1})", f"=SUM(D4:D{tr - 1})",
                       f"=SUM(E4:E{tr - 1})", f"=IF(B{tr}=0,0,C{tr}/B{tr})"], 1):
    cell = ut.cell(row=tr, column=c, value=v)
    cell.font, cell.border, cell.alignment = font(bold=True), BOX, CENTER
ut.cell(row=tr, column=6).number_format = "0.0%"
kr = tr + 1
ut.cell(row=kr, column=1, value="Skenario kritis Tidak Sesuai").font = font(bold=True)
ut.cell(row=kr, column=2, value=f'=COUNTIFS({krit},"Ya",{res},"Tidak Sesuai")').font = font(bold=True)
ut.cell(row=kr + 1, column=1, value="KESIMPULAN").font = font(bold=True, size=11)
ut.cell(row=kr + 1, column=2, value=(
    f'=IF(E{tr}>0,"Pengujian belum lengkap ("&E{tr}&" skenario belum diuji)",'
    f'IF(OR(F{tr}<0.9,B{kr}>0),"Belum diterima",IF(D{tr}>0,"Diterima dengan perbaikan","Diterima")))'))
ut.cell(row=kr + 1, column=2).font = font(bold=True, size=11, color="0B5C39")

head(ut, S0 - 1, ["Kode", "Kelompok", "Skenario", "Kritis", "Hasil", "Catatan"])
for k, (it, g) in enumerate(scen):
    r = S0 + k
    vals = [it["id"], f'{g["code"]}. {g["title"]}', it["name"], "Ya" if it.get("critical") else "–", None, None]
    for c, v in enumerate(vals, 1):
        cell = ut.cell(row=r, column=c, value=v)
        cell.border, cell.font = BOX, font()
        cell.alignment = CENTER if c in (1, 4, 5) else WRAP
    ut.cell(row=r, column=5).fill = INPUT
    ut.cell(row=r, column=6).fill = INPUT
dv_res = DataValidation(type="list", formula1='"Sesuai,Tidak Sesuai"', allow_blank=True)
ut.add_data_validation(dv_res)
dv_res.add(f"E{S0}:E{S1}")
ut.conditional_formatting.add(f"E{S0}:E{S1}", CellIsRule(operator="equal", formula=['"Sesuai"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ut.conditional_formatting.add(f"E{S0}:E{S1}", CellIsRule(operator="equal", formula=['"Tidak Sesuai"'], fill=PatternFill("solid", fgColor="FFC7CE")))
ut.freeze_panes = f"A{S0}"

wb.save(OUT)
print("OK", OUT, f"{len(scen)} skenario UAT, {N_RESP} baris responden SUS")
